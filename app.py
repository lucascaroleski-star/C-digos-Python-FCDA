import streamlit as st
import docx
import pandas as pd
from fpdf import FPDF
import re
import io
import zipfile

# ==============================================================================
# 1. CONFIGURAÇÃO DA PÁGINA
# ==============================================================================
st.set_page_config(page_title="Gerador de FCDA em PDF", layout="centered")
st.title("🚁 Gerador de FCDA (Exportação em PDF)")
st.markdown("Faça o upload da FADT (Word) para extrair os dados e gerar os formulários FCDA em PDF.")

# ==============================================================================
# 2. BANCO DE DADOS E FUNÇÕES DE LIMPEZA
# ==============================================================================
bd_aeronaves = {
    "PRHTV": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRCLR": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHEM": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHAS": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHSR": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PREBH": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PSHRM": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHFI": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PPPIT": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRMGJ": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PTHZF": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRVCA": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRBII": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHCT": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PPMIG": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PTHZS": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHLL": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PPJJJ": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHSC": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHLU": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PSHAM": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHME": {"Modelo": "AS 350 B2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PSHSL": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PSHBM": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PSHCI": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRYIT": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PSHCL": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHAE": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PPHZB": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PSGEA": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHGL": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHNB": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PSHCF": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PSHCH": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PSHCM": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PSBDF": {"Modelo": "AS 350 B3", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRHBZ": {"Modelo": "EC 130 B4", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRCBH": {"Modelo": "EC 130 B4", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRECB": {"Modelo": "EC 130 B4", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRBOP": {"Modelo": "EC 130 B4", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRDHL": {"Modelo": "EC 130 B4", "Fabricante": "AIRBUS HELICOPTERS"},
    "PSHTT": {"Modelo": "EC 130 T2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PPPLD": {"Modelo": "EC 130 T2", "Fabricante": "AIRBUS HELICOPTERS"},
    "PRFPM": {"Modelo": "EC 120B",  "Fabricante": "AIRBUS HELICOPTERS"},
    "PRRFA": {"Modelo": "EC 135 T2+","Fabricante": "AIRBUS HELICOPTERS"},
    "PRRFC": {"Modelo": "EC 135 T2+","Fabricante": "AIRBUS HELICOPTERS"},
    "PSHCC": {"Modelo": "EC 135 T1",  "Fabricante": "AIRBUS HELICOPTERS"},
    "PSHCE": {"Modelo": "EC 135 T2",  "Fabricante": "AIRBUS HELICOPTERS"},
    "PRCBM": {"Modelo": "EC 135 T2",  "Fabricante": "AIRBUS HELICOPTERS"},
    "PPEHY": {"Modelo": "HB350B",   "Fabricante": "AIRBUS HELICOPTERS"},
    "PPEHZ": {"Modelo": "HB350B",   "Fabricante": "AIRBUS HELICOPTERS"},
    "PTHTC": {"Modelo": "206B",     "Fabricante": "BELL HELICOPTER"},
    "PTHUR": {"Modelo": "206B",     "Fabricante": "BELL HELICOPTER"},
    "PRHPB": {"Modelo": "206B",     "Fabricante": "BELL HELICOPTER"},
    "PPEJI": {"Modelo": "206B",     "Fabricante": "BELL HELICOPTER"},
    "PRHSU": {"Modelo": "206L4",    "Fabricante": "BELL HELICOPTER"},
    "PTYEL": {"Modelo": "206L4",    "Fabricante": "BELL HELICOPTER"},
    "PRHIB": {"Modelo": "206L4",    "Fabricante": "BELL HELICOPTER"},
    "PRHMA": {"Modelo": "206L4",    "Fabricante": "BELL HELICOPTER"},
    "PTYUQ": {"Modelo": "407",      "Fabricante": "BELL HELICOPTER"},
    "PRGDF": {"Modelo": "407",      "Fabricante": "BELL HELICOPTER"},
    "PPJRX": {"Modelo": "505",      "Fabricante": "BELL HELICOPTER"},
    "PPJRH": {"Modelo": "505",      "Fabricante": "BELL HELICOPTER"},
    "PSPFG": {"Modelo": "412EP",    "Fabricante": "BELL HELICOPTER"},
    "PSPFH": {"Modelo": "412EP",    "Fabricante": "BELL HELICOPTER"},
    "PSFGR": {"Modelo": "429",      "Fabricante": "BELL HELICOPTER"},
    "PSHAF": {"Modelo": "429",      "Fabricante": "BELL HELICOPTER"},
    "PRFMS": {"Modelo": "429",      "Fabricante": "BELL HELICOPTER"},
    "PSSFS": {"Modelo": "429",      "Fabricante": "BELL HELICOPTER"},
    "PREFB": {"Modelo": "A109A",    "Fabricante": "LEONARDO"},
    "PREBZ": {"Modelo": "A109K2",   "Fabricante": "LEONARDO"},
    "PSIBA": {"Modelo": "AW119MKII","Fabricante": "LEONARDO"},
    "PSIBB": {"Modelo": "AW119MKII","Fabricante": "LEONARDO"},
    "PSIBC": {"Modelo": "AW119MKII","Fabricante": "LEONARDO"},
    "PSIBD": {"Modelo": "AW119MKII","Fabricante": "LEONARDO"},
    "PSIBE": {"Modelo": "AW119MKII","Fabricante": "LEONARDO"},
    "PSIBF": {"Modelo": "AW119MKII","Fabricante": "LEONARDO"},
    "PSIBG": {"Modelo": "AW119MKII","Fabricante": "LEONARDO"},
    "PSIBH": {"Modelo": "AW119MKII","Fabricante": "LEONARDO"},
    "PSIBI": {"Modelo": "AW119MKII","Fabricante": "LEONARDO"},
    "PPADN": {"Modelo": "R44II",    "Fabricante": "ROBINSON"},
    "PRYFH": {"Modelo": "R44II",    "Fabricante": "ROBINSON"},
    "PPPRL": {"Modelo": "R44II",    "Fabricante": "ROBINSON"},
    "PSHPM": {"Modelo": "R66",      "Fabricante": "ROBINSON"},
    "PSHPR": {"Modelo": "R66",      "Fabricante": "ROBINSON"},
    "PSEPH": {"Modelo": "R66",      "Fabricante": "ROBINSON"},
    "PSHLT": {"Modelo": "R66",      "Fabricante": "ROBINSON"},
    "PSHCJ": {"Modelo": "R66",      "Fabricante": "ROBINSON"}
}

@st.cache_data
def interpretar_matricula(texto_cru):
    match = re.search(r'([A-Z]{2}-?[A-Z]{3})\s*(?:\(S/?N\s*([A-Z0-9]+)\))?', texto_cru.upper())
    if match:
        mat_com_hifen = match.group(1)
        if "-" not in mat_com_hifen and len(mat_com_hifen) == 5:
            mat_com_hifen = mat_com_hifen[:2] + "-" + mat_com_hifen[2:]
        mat_limpa = match.group(1).replace("-", "")
        sn = match.group(2) if match.group(2) else "N/A"
    else:
        mat_com_hifen = str(texto_cru).strip()
        mat_limpa = mat_com_hifen.replace("-", "").replace(" ", "")
        sn = "N/A"

    dados_bd = bd_aeronaves.get(mat_limpa, {"Modelo": "Verificar", "Fabricante": "Verificar"})
    return mat_com_hifen, sn, dados_bd["Modelo"], dados_bd["Fabricante"]

def obter_texto_com_checkbox(cell):
    text_parts = []
    for element in cell._element.iter():
        tag = element.tag.split('}')[-1].lower()
        if tag == 't' and element.text:
            text_parts.append(element.text)
        elif tag == 'tab':
            text_parts.append('\t')
        elif tag == 'sym':
            char_val = next((element.get(k) for k in element.keys() if k.endswith('char')), None)
            if char_val:
                char_val = char_val.upper()
                if char_val in ['F0FE', 'F0FD', 'F058', '00FE', '00FD', '0058', '2611']: text_parts.append('[X] ')
                elif char_val in ['F0A8', 'F0A1', '00A8', '00A1', '2610']: text_parts.append('[ ] ')
        elif tag == 'checkbox':
            checked = False
            for child in element.iter():
                if child.tag.split('}')[-1].lower() == 'checked':
                    val = next((child.get(k) for k in child.keys() if k.endswith('val')), None)
                    checked = val in ['1', 'true', 'True'] or val is None
            text_parts.append('[X] ' if checked else '[ ] ')

    full_text = "".join(text_parts).replace('☒', '[X]').replace('☐', '[ ]').replace('☑', '[X]')
    full_text = full_text.replace('\n', ' ')
    full_text = re.sub(r'(\[X\]\s*)+', '[X] ', full_text)
    full_text = re.sub(r'(\[\s\]\s*)+', '[ ] ', full_text)
    return re.sub(r' +', ' ', full_text).strip()

# ==============================================================================
# 3. FUNÇÃO QUE DESENHA O PDF
# ==============================================================================
def criar_pdf_fcda(dados):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    # Resolve caracteres especiais (evita erro no FPDF com acentos)
    def clean_text(text):
        return str(text).encode('latin-1', 'replace').decode('latin-1')

    # Título do Documento
    pdf.set_font("Helvetica", 'B', 16)
    pdf.cell(0, 10, clean_text(f"FCDA - MATRÍCULA: {dados['matricula']}"), ln=True, align='C')
    pdf.ln(5)

    # Função auxiliar para criar as "linhas" do formulário
    def adicionar_campo(titulo, valor):
        pdf.set_font("Helvetica", 'B', 11)
        pdf.cell(0, 6, clean_text(titulo), ln=True)
        pdf.set_font("Helvetica", '', 11)
        pdf.multi_cell(0, 6, clean_text(valor))
        pdf.ln(3)

    adicionar_campo("1 - Aeronave/Modelo:", dados['matricula'])
    adicionar_campo("2 - Diretriz:", dados['diretriz'])
    adicionar_campo("3 - Efetividade:", dados['data_efetiva'])
    adicionar_campo("4 - Vencimento(s):", dados['vencimento'])
    adicionar_campo("5 - Aplicável:", dados['aplicavel_tipo'])
    adicionar_campo("6 - Tipo de ação:", dados['tipo_acao'])
    adicionar_campo("7 - Análise:", dados['analise'])
    adicionar_campo("8 - Justificativa:", dados['justificativa'])
    adicionar_campo("9/10 - Documentos de Referência:", dados['docs_ref'])

    pdf.ln(5)
    pdf.set_font("Helvetica", 'B', 12)
    pdf.cell(0, 8, clean_text("11 - Identificação do Produto Aeronáutico"), ln=True, align='C')
    pdf.line(10, pdf.get_y(), 200, pdf.get_y()) # Linha divisória
    pdf.ln(2)
    
    adicionar_campo("Fabricante:", dados['fabricante'])
    adicionar_campo("Modelo:", dados['modelo'])
    adicionar_campo("S/N:", dados['sn'])

    # Retorna o PDF em formato de bytes para salvar no ZIP
    return pdf.output(dest='S')

# ==============================================================================
# 4. INTERFACE E EXTRAÇÃO
# ==============================================================================
fadt_file = st.file_uploader("📄 Anexar FADT Analisada (.docx)", type=["docx"])

if fadt_file:
    if st.button("🚀 Gerar FCDAs em PDF (ZIP)", use_container_width=True):
        with st.spinner('Extraindo dados e desenhando PDFs...'):
            
            # --- LEITURA DO WORD ---
            doc = docx.Document(fadt_file)
            dados_cabecalho, aplicaveis, nao_aplicaveis = [], [], []
            modo_atual = 'cabecalho'

            for table in doc.tables:
                for row in table.rows:
                    celulas_unicas, seen_cells = [], set()
                    for cell in row.cells:
                        if cell._element not in seen_cells:
                            seen_cells.add(cell._element)
                            texto_celula = obter_texto_com_checkbox(cell)
                            if texto_celula: celulas_unicas.append(texto_celula)

                    if len(celulas_unicas) == 1 and '\t' in celulas_unicas[0]:
                        celulas_unicas = [t.strip() for t in celulas_unicas[0].split('\t') if t.strip()]
                    row_text = " ".join(celulas_unicas).strip()

                    if "Motivo da emissão deste documento" in row_text: modo_atual = None; continue
                    elif "APLICÁVEL" in row_text and "NÃO APLICÁVEL" not in row_text: modo_atual = 'aplicavel'; continue
                    elif "NÃO APLICÁVEL" in row_text: modo_atual = 'nao_aplicavel'; continue
                    elif "- - - Fim da Lista" in row_text: modo_atual = None; continue

                    if not celulas_unicas: continue

                    if modo_atual == 'cabecalho': dados_cabecalho.append(celulas_unicas)
                    elif modo_atual == 'aplicavel':
                        if "Matrícula" in row_text or "Limitação" in row_text: continue
                        aplicaveis.append([
                            celulas_unicas[0] if len(celulas_unicas) > 0 else "",
                            celulas_unicas[1] if len(celulas_unicas) > 1 else "",
                            celulas_unicas[2] if len(celulas_unicas) > 2 else ""
                        ])
                    elif modo_atual == 'nao_aplicavel':
                        if "Matrícula" in row_text or "Justificativa" in row_text: continue
                        nao_aplicaveis.append([
                            celulas_unicas[0] if len(celulas_unicas) > 0 else "",
                            celulas_unicas[1] if len(celulas_unicas) > 1 else ""
                        ])

            # --- ANÁLISE DO CABEÇALHO ---
            texto_cabecalho = " ".join([" ".join(row) for row in dados_cabecalho])
            match_ad = re.search(r'(AD\s\d{4}-\d{4})', texto_cabecalho)
            ext_diretriz = match_ad.group(1) if match_ad else "Verificar no AD"
            datas = re.findall(r'\d{2}/\d{2}/\d{4}', texto_cabecalho)
            ext_data_efetiva = datas[1] if len(datas) >= 2 else (datas[0] if datas else "Verificar")
            match_docs = re.search(r'ASB[\w\s,-]+', texto_cabecalho)
            ext_docs = match_docs.group(0).strip() if match_docs else "Verificar Documentos"
            ext_tipo = "Aeronave" if "[X] Aeronave" in texto_cabecalho else ("Motor" if "[X] Motor" in texto_cabecalho else "Verificar")
            ext_acao = "Terminal" if "[X] Terminal" in texto_cabecalho else ("Repetitiva" if "[X] Repetitiva" in texto_cabecalho else "Verificar")

            # --- GERAÇÃO DOS PDFs E EMPACOTAMENTO ZIP ---
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
                
                def injetar_e_salvar_pdf(mat_crua, vencimento, analise, justificativa):
                    mat_hifen, sn, modelo, fab = interpretar_matricula(mat_crua)
                    
                    dados = {
                        'matricula': mat_hifen, 'diretriz': ext_diretriz, 'data_efetiva': ext_data_efetiva,
                        'vencimento': vencimento, 'aplicavel_tipo': ext_tipo, 'tipo_acao': ext_acao,
                        'analise': analise, 'justificativa': justificativa, 'docs_ref': ext_docs,
                        'fabricante': fab, 'modelo': modelo, 'sn': sn
                    }
                    
                    pdf_bytes = criar_pdf_fcda(dados)
                    # Salva direto na raiz do ZIP, sem pastas
                    zip_file.writestr(f"FCDA_{mat_hifen}.pdf", pdf_bytes)

                for a in aplicaveis: injetar_e_salvar_pdf(a[0], a[1], 'SIM', a[2])
                for na in nao_aplicaveis: injetar_e_salvar_pdf(na[0], 'N/A', 'NÃO', na[1])

            # --- BOTÃO DE DOWNLOAD FINAL ---
            st.success("✅ Todos os PDFs gerados com sucesso!")
            st.download_button(
                label="📥 Baixar todos os FCDAs (Formato PDF)",
                data=zip_buffer.getvalue(),
                file_name="Lote_FCDAs_PDF.zip",
                mime="application/zip",
                use_container_width=True
            )
